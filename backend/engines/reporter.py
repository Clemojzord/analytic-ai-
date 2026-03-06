"""
Analytic AI — Report Engine
PDF (ReportLab) + Excel (openpyxl) report generation
"""
from __future__ import annotations

import io
import json
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (
    HRFlowable, Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
)
from reportlab.lib.enums import TA_CENTER, TA_LEFT
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ── Brand colors ─────────────────────────────────────────────
BRAND_PRIMARY   = colors.HexColor("#38bdf8")
BRAND_DARK      = colors.HexColor("#0d1424")
BRAND_LIGHT     = colors.HexColor("#f0f9ff")
BRAND_MUTED     = colors.HexColor("#94a3b8")
BRAND_SUCCESS   = colors.HexColor("#34d399")
BRAND_WARNING   = colors.HexColor("#fbbf24")
BRAND_DANGER    = colors.HexColor("#fb7185")


# ─────────────────────────────────────────────────────────────
# Public API
# ─────────────────────────────────────────────────────────────

def generate_pdf(
    dataset_name: str,
    kpis: dict[str, Any],
    insights: dict[str, Any],
    output_path: Path,
    df_preview: pd.DataFrame | None = None,
) -> Path:
    """Generate a branded PDF report and save to output_path."""
    output_path.parent.mkdir(parents=True, exist_ok=True)

    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=A4,
        rightMargin=2 * cm,
        leftMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
        title=f"Analytic AI Report — {dataset_name}",
    )

    styles = _get_styles()
    story: list = []

    # ── Cover ─────────────────────────────────────────────────
    story.append(Spacer(1, 1.5 * cm))
    story.append(Paragraph("📊 Analytic AI", styles["title"]))
    story.append(Paragraph(f"Business Intelligence Report", styles["subtitle"]))
    story.append(Paragraph(f"Dataset: {dataset_name}", styles["subtitle_sm"]))
    story.append(Paragraph(
        f"Generated: {datetime.utcnow().strftime('%d %B %Y, %H:%M UTC')}",
        styles["muted"]
    ))
    story.append(HRFlowable(width="100%", thickness=1, color=BRAND_PRIMARY, spaceAfter=20))

    # ── Executive Summary ─────────────────────────────────────
    story.append(Paragraph("Executive Summary", styles["section"]))
    summary_text = insights.get("executive_summary", "No summary available.")
    # Strip markdown bold markers for PDF
    summary_text = summary_text.replace("**", "").replace("⬆", "↑").replace("⬇", "↓")
    for line in summary_text.split("\n"):
        if line.strip():
            story.append(Paragraph(line.strip(), styles["body"]))
            story.append(Spacer(1, 4))
    story.append(Spacer(1, 0.5 * cm))

    # ── KPI Table ─────────────────────────────────────────────
    story.append(Paragraph("Key Performance Indicators", styles["section"]))
    kpi_rows = _build_kpi_table_data(kpis)
    if kpi_rows:
        tbl = Table(kpi_rows, colWidths=[9 * cm, 7 * cm])
        tbl.setStyle(_kpi_table_style())
        story.append(tbl)
        story.append(Spacer(1, 0.5 * cm))

    # ── Monthly Revenue ───────────────────────────────────────
    monthly = kpis.get("monthly_revenue", [])
    if monthly:
        story.append(Paragraph("Monthly Revenue Breakdown", styles["section"]))
        mdata = [["Period", "Revenue (KES)", "Growth %"]]
        for m in monthly:
            growth_str = f"{m['growth_rate']:+.1f}%" if m.get("growth_rate") is not None else "—"
            mdata.append([m["period"], f"KES {m['revenue']:,.2f}", growth_str])
        mtbl = Table(mdata, colWidths=[5 * cm, 7 * cm, 4 * cm])
        mtbl.setStyle(_monthly_table_style())
        story.append(mtbl)
        story.append(Spacer(1, 0.5 * cm))

    # ── Insights ──────────────────────────────────────────────
    story.append(PageBreak())
    story.append(Paragraph("Automated Insights", styles["section"]))
    for item in insights.get("insights", []):
        sev = item.get("severity", "info")
        icon = {"critical": "⛔", "warning": "⚠️", "info": "ℹ️"}.get(sev, "•")
        story.append(Paragraph(
            f"{icon}  {item['text']}",
            styles["insight_" + sev]
        ))
        story.append(Spacer(1, 6))

    # ── Data Preview ──────────────────────────────────────────
    if df_preview is not None and not df_preview.empty:
        story.append(Spacer(1, 0.5 * cm))
        story.append(Paragraph("Data Preview (first 10 rows)", styles["section"]))
        preview = df_preview.head(10).astype(str)
        cols = [c[:20] for c in preview.columns.tolist()]
        pdata = [cols] + preview.values.tolist()
        col_w = [16 * cm / max(len(cols), 1)] * len(cols)
        ptbl = Table(pdata, colWidths=col_w, repeatRows=1)
        ptbl.setStyle(_preview_table_style())
        story.append(ptbl)

    # ── Footer ────────────────────────────────────────────────
    story.append(Spacer(1, 1 * cm))
    story.append(HRFlowable(width="100%", thickness=0.5, color=BRAND_MUTED))
    story.append(Paragraph(
        "Generated by Analytic AI · analytic-ai-icg.web.app · © 2026",
        styles["footer"]
    ))

    doc.build(story)
    return output_path


def generate_excel(
    dataset_name: str,
    kpis: dict[str, Any],
    insights: dict[str, Any],
    output_path: Path,
    df_cleaned: pd.DataFrame | None = None,
) -> Path:
    """Generate a multi-sheet Excel report."""
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()

    # Sheet 1: KPI Summary
    ws_kpi = wb.active
    ws_kpi.title = "KPI Summary"
    _write_kpi_sheet(ws_kpi, dataset_name, kpis)

    # Sheet 2: Monthly Revenue
    ws_monthly = wb.create_sheet("Monthly Revenue")
    _write_monthly_sheet(ws_monthly, kpis.get("monthly_revenue", []))

    # Sheet 3: Insights
    ws_insights = wb.create_sheet("Insights")
    _write_insights_sheet(ws_insights, insights)

    # Sheet 4: Cleaned Data
    if df_cleaned is not None and not df_cleaned.empty:
        ws_data = wb.create_sheet("Cleaned Data")
        _write_data_sheet(ws_data, df_cleaned)

    wb.save(str(output_path))
    return output_path


# ─────────────────────────────────────────────────────────────
# PDF Helpers
# ─────────────────────────────────────────────────────────────

def _get_styles() -> dict:
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("title", parent=base["Title"],
                                fontSize=26, textColor=BRAND_PRIMARY,
                                spaceAfter=6, fontName="Helvetica-Bold"),
        "subtitle": ParagraphStyle("subtitle", parent=base["Normal"],
                                   fontSize=13, textColor=BRAND_MUTED,
                                   spaceAfter=4, alignment=TA_LEFT),
        "subtitle_sm": ParagraphStyle("subtitle_sm", parent=base["Normal"],
                                      fontSize=11, textColor=BRAND_MUTED,
                                      spaceAfter=4),
        "section": ParagraphStyle("section", parent=base["Heading2"],
                                  fontSize=14, textColor=BRAND_PRIMARY,
                                  fontName="Helvetica-Bold",
                                  spaceBefore=14, spaceAfter=8),
        "body": ParagraphStyle("body", parent=base["Normal"],
                               fontSize=10, textColor=colors.HexColor("#334155"),
                               leading=16),
        "muted": ParagraphStyle("muted", parent=base["Normal"],
                                fontSize=9, textColor=BRAND_MUTED, spaceAfter=10),
        "insight_info": ParagraphStyle("insight_info", parent=base["Normal"],
                                       fontSize=9.5, textColor=colors.HexColor("#1e40af"),
                                       backColor=colors.HexColor("#eff6ff"),
                                       leftIndent=8, rightIndent=8,
                                       borderPadding=4, leading=14),
        "insight_warning": ParagraphStyle("insight_warning", parent=base["Normal"],
                                          fontSize=9.5, textColor=colors.HexColor("#92400e"),
                                          backColor=colors.HexColor("#fffbeb"),
                                          leftIndent=8, rightIndent=8,
                                          borderPadding=4, leading=14),
        "insight_critical": ParagraphStyle("insight_critical", parent=base["Normal"],
                                           fontSize=9.5, textColor=colors.HexColor("#991b1b"),
                                           backColor=colors.HexColor("#fff1f2"),
                                           leftIndent=8, rightIndent=8,
                                           borderPadding=4, leading=14),
        "footer": ParagraphStyle("footer", parent=base["Normal"],
                                 fontSize=7.5, textColor=BRAND_MUTED,
                                 alignment=TA_CENTER, spaceBefore=4),
    }


def _build_kpi_table_data(kpis: dict) -> list:
    label_map = {
        "total_revenue":         "Total Revenue (KES)",
        "net_profit":            "Net Profit (KES)",
        "profit_margin_pct":     "Profit Margin (%)",
        "mom_growth_rate_pct":   "Month-on-Month Growth (%)",
        "avg_monthly_growth_pct":"Avg Monthly Growth (%)",
        "top_product":           "Best Selling Product",
        "top_branch":            "Top Performing Branch",
        "avg_transaction_value": "Avg Transaction Value (KES)",
        "unique_clients":        "Unique Customers",
        "estimated_clv_annual":  "Customer LTV — Annual (KES)",
        "inventory_turnover":    "Inventory Turnover (×)",
        "total_records":         "Total Transactions",
    }
    rows = [["KPI", "Value"]]
    for key, label in label_map.items():
        val = kpis.get(key)
        if val is None:
            continue
        if isinstance(val, float):
            formatted = f"{val:,.2f}"
        else:
            formatted = str(val)
        rows.append([label, formatted])
    return rows


def _kpi_table_style() -> TableStyle:
    return TableStyle([
        ("BACKGROUND",    (0, 0), (-1, 0), BRAND_PRIMARY),
        ("TEXTCOLOR",     (0, 0), (-1, 0), colors.white),
        ("FONTNAME",      (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE",      (0, 0), (-1, 0), 10),
        ("ROWBACKGROUNDS",(0, 1), (-1, -1), [colors.HexColor("#f8fafc"), colors.white]),
        ("FONTSIZE",      (0, 1), (-1, -1), 9.5),
        ("GRID",          (0, 0), (-1, -1), 0.4, colors.HexColor("#e2e8f0")),
        ("LEFTPADDING",   (0, 0), (-1, -1), 10),
        ("RIGHTPADDING",  (0, 0), (-1, -1), 10),
        ("TOPPADDING",    (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ])


def _monthly_table_style() -> TableStyle:
    return TableStyle([
        ("BACKGROUND",  (0, 0), (-1, 0), colors.HexColor("#1e293b")),
        ("TEXTCOLOR",   (0, 0), (-1, 0), BRAND_LIGHT),
        ("FONTNAME",    (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#f0f9ff"), colors.white]),
        ("FONTSIZE",    (0, 0), (-1, -1), 9),
        ("GRID",        (0, 0), (-1, -1), 0.3, colors.HexColor("#e2e8f0")),
        ("ALIGN",       (1, 0), (-1, -1), "RIGHT"),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING",  (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING",(0, 0), (-1, -1), 5),
    ])


def _preview_table_style() -> TableStyle:
    return TableStyle([
        ("BACKGROUND",  (0, 0), (-1, 0), colors.HexColor("#0f172a")),
        ("TEXTCOLOR",   (0, 0), (-1, 0), BRAND_LIGHT),
        ("FONTNAME",    (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE",    (0, 0), (-1, -1), 7),
        ("GRID",        (0, 0), (-1, -1), 0.3, colors.HexColor("#e2e8f0")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#f8fafc"), colors.white]),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING",  (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING",(0, 0), (-1, -1), 3),
    ])


# ─────────────────────────────────────────────────────────────
# Excel Helpers
# ─────────────────────────────────────────────────────────────

def _xl_header_style():
    return {
        "font":      Font(bold=True, color="FFFFFF", size=11),
        "fill":      PatternFill("solid", fgColor="0F172A"),
        "alignment": Alignment(horizontal="center", vertical="center"),
        "border": Border(
            bottom=Side(style="thin", color="38BDF8"),
        ),
    }


def _apply_header(cell):
    s = _xl_header_style()
    cell.font = s["font"]
    cell.fill = s["fill"]
    cell.alignment = s["alignment"]
    cell.border = s["border"]


def _write_kpi_sheet(ws, name: str, kpis: dict) -> None:
    ws.column_dimensions["A"].width = 35
    ws.column_dimensions["B"].width = 20

    ws["A1"] = "📊 Analytic AI — KPI Summary"
    ws["A1"].font = Font(bold=True, size=16, color="38BDF8")
    ws["A2"] = f"Dataset: {name}"
    ws["A2"].font = Font(italic=True, color="94A3B8")
    ws["A3"] = f"Generated: {datetime.utcnow().strftime('%d %B %Y, %H:%M UTC')}"
    ws["A3"].font = Font(italic=True, color="94A3B8", size=9)

    ws.append([])
    header_row = ws.max_row + 1
    ws.append(["KPI", "Value"])
    for cell in ws[header_row]:
        _apply_header(cell)

    kpis_display = {
        "total_revenue":         ("Total Revenue (KES)", lambda v: f"{v:,.2f}"),
        "net_profit":            ("Net Profit (KES)", lambda v: f"{v:,.2f}"),
        "profit_margin_pct":     ("Profit Margin (%)", lambda v: f"{v:.2f}%"),
        "mom_growth_rate_pct":   ("Month-on-Month Growth (%)", lambda v: f"{v:+.2f}%"),
        "top_product":           ("Best Selling Product", str),
        "top_branch":            ("Top Branch", str),
        "avg_transaction_value": ("Avg Transaction Value (KES)", lambda v: f"{v:,.2f}"),
        "unique_clients":        ("Unique Customers", str),
        "estimated_clv_annual":  ("Annual Customer LTV (KES)", lambda v: f"{v:,.2f}"),
        "inventory_turnover":    ("Inventory Turnover (×)", lambda v: f"{v:.2f}"),
        "total_records":         ("Total Transactions", str),
    }
    for key, (label, fmt) in kpis_display.items():
        val = kpis.get(key)
        if val is not None:
            ws.append([label, fmt(val)])


def _write_monthly_sheet(ws, monthly: list) -> None:
    ws.column_dimensions["A"].width = 15
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 15

    ws.append(["Period", "Revenue (KES)", "Growth %"])
    for cell in ws[1]:
        _apply_header(cell)

    for m in monthly:
        ws.append([
            m.get("period", ""),
            m.get("revenue", 0),
            m.get("growth_rate") or "",
        ])


def _write_insights_sheet(ws, insights: dict) -> None:
    ws.column_dimensions["A"].width = 15
    ws.column_dimensions["B"].width = 80
    ws.column_dimensions["C"].width = 15

    ws.append(["Category", "Insight", "Severity"])
    for cell in ws[1]:
        _apply_header(cell)

    severity_colors = {"critical": "FEF2F2", "warning": "FFFBEB", "info": "EFF6FF"}
    for item in insights.get("insights", []):
        row = [item.get("category", ""), item.get("text", ""), item.get("severity", "")]
        ws.append(row)
        sev = item.get("severity", "info")
        fill_color = severity_colors.get(sev, "FFFFFF")
        for cell in ws[ws.max_row]:
            cell.fill = PatternFill("solid", fgColor=fill_color)
            cell.alignment = Alignment(wrap_text=True)

    # Executive summary in last rows
    ws.append([])
    ws.append(["Executive Summary", "", ""])
    for line in insights.get("executive_summary", "").split("\n"):
        if line.strip():
            ws.append(["", line.replace("**", "").replace("📊", "").strip(), ""])


def _write_data_sheet(ws, df: pd.DataFrame) -> None:
    # Drop internal outlier flag columns for export
    export_df = df[[c for c in df.columns if not c.startswith("_outlier_")]]
    cols = list(export_df.columns)
    ws.append(cols)
    for cell in ws[1]:
        _apply_header(cell)

    for row in export_df.itertuples(index=False):
        ws.append(list(row))

    # Auto-width
    for i, col in enumerate(cols, 1):
        max_len = max(len(str(col)), 12)
        ws.column_dimensions[get_column_letter(i)].width = min(max_len + 4, 40)
